# -*- coding: utf-8 -*-
import os
import docx
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

TEMPLATE_PATH = r"C:\Users\user\.gemini\antigravity\brain\e80facd2-93fc-43d0-816c-5cfbe331621f\.user_uploaded\media_1793019983273_dc98dac1.docx"
OUT_DIR = r"C:\Users\user\.gemini\antigravity\scratch\cyber_crypto_labs\reports"


def set_run_font(run, name="Times New Roman", size_pt=14, bold=False, italic=False, color=None):
    run.font.name = name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor(*color)


def format_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_cm=1.25, space_after_pt=0, space_before_pt=0, line_spacing=1.5):
    p.alignment = align
    pf = p.paragraph_format
    pf.first_line_indent = Cm(first_line_cm) if first_line_cm else Cm(0)
    pf.space_after = Pt(space_after_pt)
    pf.space_before = Pt(space_before_pt)
    pf.line_spacing = line_spacing


def prepare_base_doc(lab_num: int, course_name: str, topic_name: str):
    doc = docx.Document(TEMPLATE_PATH)

    # Remove all tables from template body
    for tbl in list(doc.tables):
        tbl._element.getparent().remove(tbl._element)

    # Keep only paragraphs 0..11 (the exact Title Page with Lviv Polytechnic crest)
    for p in list(doc.paragraphs[12:]):
        p._element.getparent().remove(p._element)

    # Update Paragraph 2: Lab number
    p2 = doc.paragraphs[2]
    p2.text = ""
    r2 = p2.add_run(f"про виконання лабораторної роботи №{lab_num}")
    set_run_font(r2, italic=True)

    # Update Paragraph 3: Course name
    p3 = doc.paragraphs[3]
    p3.text = ""
    r3 = p3.add_run(f"з курсу «{course_name}»")
    set_run_font(r3, italic=True)

    # Update Paragraph 4: Topic
    p4 = doc.paragraphs[4]
    p4.text = ""
    r4 = p4.add_run(f"на тему: «{topic_name}»")
    set_run_font(r4, italic=True)

    # Add page break at end of paragraph 11 ("Львів 2026") so content starts cleanly on page 2
    doc.paragraphs[11].runs[-1].add_break(docx.enum.text.WD_BREAK.PAGE)

    return doc


def add_meta_paragraph(doc, meta_text: str):
    p = doc.add_paragraph()
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_cm=1.25)
    r_bold = p.add_run("Мета роботи - ")
    set_run_font(r_bold, bold=True)
    r_text = p.add_run(meta_text)
    set_run_font(r_text, bold=False)


def add_heading_line(doc, text: str, centered=False):
    p = doc.add_paragraph()
    format_paragraph(
        p,
        align=WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT,
        first_line_cm=0 if centered else 1.25,
        space_before_pt=6,
        space_after_pt=4
    )
    r = p.add_run(text)
    set_run_font(r, bold=True)


def add_body_paragraph(doc, text: str, bold_prefix: str = None):
    p = doc.add_paragraph()
    format_paragraph(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line_cm=1.25)
    if bold_prefix:
        rb = p.add_run(bold_prefix)
        set_run_font(rb, bold=True)
    rt = p.add_run(text)
    set_run_font(rt, bold=False)


def add_code_block(doc, code_str: str):
    for line in code_str.strip().split("\n"):
        p = doc.add_paragraph()
        format_paragraph(
            p,
            align=WD_ALIGN_PARAGRAPH.LEFT,
            first_line_cm=0.5,
            space_after_pt=0,
            space_before_pt=0,
            line_spacing=1.15
        )
        r = p.add_run(line)
        set_run_font(r, name="Consolas", size_pt=10.5, bold=False)


def add_screenshot_placeholder(doc, instruction_text: str, caption_text: str):
    p_box = doc.add_paragraph()
    format_paragraph(p_box, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_cm=0, space_before_pt=8, space_after_pt=4)
    r_box = p_box.add_run(f"[ВСТАВИТИ СКРІНШОТ ТУТ: {instruction_text}]")
    set_run_font(r_box, bold=True, italic=True, color=(30, 64, 175))

    p_cap = doc.add_paragraph()
    format_paragraph(p_cap, align=WD_ALIGN_PARAGRAPH.CENTER, first_line_cm=0, space_before_pt=2, space_after_pt=8)
    r_cap = p_cap.add_run(caption_text)
    set_run_font(r_cap, bold=False)


def build_lab1():
    doc = prepare_base_doc(
        lab_num=1,
        course_name="Криптографія та захист інформації",
        topic_name="ШИФР ЦЕЗАРЯ. ЧАСТОТНИЙ КРИПТОАНАЛІЗ. ШИФР ПЕРЕСТАНОВКИ"
    )

    add_meta_paragraph(
        doc,
        "ознайомитися з класичними симетричними алгоритмами шифрування моноалфавітної заміни (шифр Цезаря) та матричної стовпцевої перестановки, дослідити статистичні закономірності природних мов і практично реалізувати програмні модулі для шифрування, дешифрування та автоматизованого частотного криптоаналізу шифротекстів."
    )

    add_heading_line(doc, "Завдання до роботи:", centered=False)
    add_body_paragraph(doc, "1. За допомогою вибраної мови програмування (Python та JavaScript у складі розробленого веб-комплексу AEGIS CRYPTO-OPS) написати програму, яка забезпечує покрокове шифрування і дешифрування та наочне виведення на екран результатів роботи шифру Цезаря для українського (33 літери) та англійського (26 літер) алфавітів.")
    add_body_paragraph(doc, "2. Реалізувати програмний модуль частотного криптоаналізу для автоматичного визначення невідомого ключа K моноалфавітного шифру на основі таблиці стандартних частот літер мови та дешифрувати надану в методичних вказівках криптограму.")
    add_body_paragraph(doc, "3. Реалізувати алгоритм шифрування та дешифрування методом стовпцевої перестановки за ключовим словом із покроковою візуалізацією матриці перестановки.")

    add_heading_line(doc, "Хід роботи", centered=True)

    add_heading_line(doc, "1. Реалізація та дослідження шифру моноалфавітної заміни (шифру Цезаря)")
    add_body_paragraph(doc, "У шифрі моноалфавітної заміни кожна літера відкритого тексту з числовою позицією m в алфавіті потужності N (де N = 26 для англійського алфавіту та N = 33 для українського алфавіту) циклічно зсувається на фіксоване секретне число K позицій вперед за формулою C = (m + K) mod N. При дешифруванні виконується обернений циклічний зсув на K позицій назад: m = (C - K) mod N.")
    add_body_paragraph(doc, "Для наочного покрокового відображення роботи алгоритму було розроблено інтерактивний веб-модуль у вкладці «ЛАБ 01 -> 1.1. Шифр Цезаря» та консольний скрипт на мові Python.")
    add_body_paragraph(doc, "Програмна реалізація функції шифрування та дешифрування шифром Цезаря:")

    add_code_block(doc, """UA_ALPHABET = "АБВГҐДЕЄЖЗИІЇЙКЛМНОПРСТУФХЦЧШЩЬЮЯ"
EN_ALPHABET = "ABCDEFGHIJKLMNOPQRSTUVWXYZ"

def caesar_cipher(text: str, key: int, decrypt: bool = False) -> str:
    alphabet = EN_ALPHABET if any(c.upper() in EN_ALPHABET for c in text) else UA_ALPHABET
    n = len(alphabet)
    shift = (-key if decrypt else key) % n
    result = []
    for ch in text:
        upper_ch = ch.upper()
        if upper_ch in alphabet:
            idx = alphabet.index(upper_ch)
            new_ch = alphabet[(idx + shift) % n]
            result.append(new_ch.lower() if ch.islower() else new_ch)
        else:
            result.append(ch)
    return "".join(result)""")

    add_body_paragraph(doc, "Перевірка роботи алгоритму на контрольному прикладі з методичних вказівок: вхідне повідомлення «HONEY, I AM HOME», ключ K = 3. У результаті покрокового зсуву кожної літери на 3 позиції праворуч (H(7)+3=10(K), O(14)+3=17(R), N(13)+3=16(Q), E(4)+3=7(H), Y(24)+3 mod 26 = 1(B)) отримано шифротекст: «KRQHB, L DP KRPH» (рис. 1).")

    add_screenshot_placeholder(
        doc,
        "Відкрийте сайт http://localhost:8090 -> вкладка «ЛАБ 01» -> підвкладка «1.1. Шифр Цезаря». Натисніть пресет «ENG: HONEY, I AM HOME (K=3)», натисніть кнопку «[Вперед]» кілька разів і зробіть скріншот екрана з криптографічним диском та стрічкою алфавіту",
        "Рисунок 1 - Інтерфейс покрокового шифрування повідомлення шифром Цезаря при K = 3"
    )

    add_body_paragraph(doc, "Також було протестовано шифрування та дешифрування тексту українською мовою для студента варіанту 12 («Марчук Дмитро Андрійович») при ключі K = 12 (рис. 2).")

    add_screenshot_placeholder(
        doc,
        "У тій самій вкладці «1.1. Шифр Цезаря» введіть у поле вхідного повідомлення «Марчук Дмитро Андрійович», виберіть ключ K = 12 та зробіть скріншот результату перетворення і таблиці відповідності індексів 0..32",
        "Рисунок 2 - Шифрування тексту українською мовою при ключі K = 12"
    )

    add_heading_line(doc, "2. Реалізація методу частотного криптоаналізу")
    add_body_paragraph(doc, "Криптоаналіз моноалфавітного шифру ґрунтується на зіставленні частоти появи символів у перехопленому шифротексті з еталонними статистичними частотами літер природної мови (в українській мові найчастішою є літера «О» з частотою 0,094 та індексом 18; в англійській мові - літера «E» з частотою 0,127 та індексом 4).")
    add_body_paragraph(doc, "Програмна реалізація автоматичного частотного криптоаналізу:")

    add_code_block(doc, """from collections import Counter

def frequency_attack_ua(ciphertext: str, most_common_lang_char: str = "О"):
    filtered = [c.upper() for c in ciphertext if c.upper() in UA_ALPHABET]
    counts = Counter(filtered)
    most_freq_cipher, freq_val = counts.most_common(1)[0]
    key = (UA_ALPHABET.index(most_freq_cipher) - UA_ALPHABET.index(most_common_lang_char)) % len(UA_ALPHABET)
    decrypted = caesar_cipher(ciphertext, key, decrypt=True)
    return key, most_freq_cipher, freq_val, decrypted""")

    add_body_paragraph(doc, "Дослідження контрольної криптограми з методичних вказівок: «ДАФИНЦШЕИЮЯЗЩШФЬИТЧИВЮЯШХСЯЯЗВИЇШЧШЮФЬСПЄСПИЇОЛРПЧИЦРЗФЬРІИЇЇШЛСЯИФСЦРЇЧЄШЦАСЇШЧШІСХЗЧИЮДАФИНЧИЮЮИЦРВЧМЦИУШЙШЧСЛМЇСЛЯШЙШЕШЇШЧШЮФЬСПЄ».")
    add_body_paragraph(doc, "За результатами автоматичного підрахунку кількості появи кожної літери встановлено, що найчастіше у криптограмі зустрічається літера «Ш» - 16 разів (позиція в українському алфавіті: 28). Ототожнюючи її з найчастішою літерою української мови «О» (позиція в алфавіті: 18), обчислюємо ключ шифрування: K = (28 - 18) mod 33 = 10. Після автоматичного дешифрування ключем K = 10 відновлено вихідний відкритий текст: «Шукаємо щастя по країнах, століттях, а воно скрізь і завжди з нами; як риба в воді, так і ми в ньому, і воно біля нас шукає нас самих. Нема його ніде від того, що воно скрізь» (рис. 3).")

    add_screenshot_placeholder(
        doc,
        "Перейдіть у підвкладку «1.2. Частотний Криптоаналіз (Авто-Злам)». Зробіть скріншот, де видно математичний висновок (K = 28 - 18 = 10), подвійну гістограму частот, Таблицю 1.4 з виділеною літерою «Ш» (16) та відновлений текст унизу",
        "Рисунок 3 - Результат роботи спектрометра частотного аналізу та автоматичне відновлення відкритого тексту"
    )

    add_heading_line(doc, "3. Реалізація шифру стовпцевої перестановки за ключовим словом")
    add_body_paragraph(doc, "У шифрі перестановки відкритий текст записується по рядках у прямокутну матрицю розмірністю R x C, де C - довжина ключового слова. Стовпці нумеруються відповідно до алфавітного порядку літер ключового слова. Криптограма формується шляхом послідовного виписування символів по стовпцях у порядку зростання їхніх номерів.")
    add_body_paragraph(doc, "Програмна реалізація шифрування та дешифрування методом перестановки:")

    add_code_block(doc, """def columnar_transposition_encrypt(plaintext: str, keyword: str) -> str:
    cols = len(keyword)
    order = sorted(range(cols), key=lambda i: (keyword[i].upper(), i))
    rows = (len(plaintext) + cols - 1) // cols
    padded = plaintext.ljust(rows * cols, " ")
    matrix = [list(padded[r * cols:(r + 1) * cols]) for r in range(rows)]
    return "".join(matrix[r][c] for c in order for r in range(rows))

def columnar_transposition_decrypt(ciphertext: str, keyword: str) -> str:
    cols = len(keyword)
    rows = (len(ciphertext) + cols - 1) // cols
    order = sorted(range(cols), key=lambda i: (keyword[i].upper(), i))
    matrix = [[""] * cols for _ in range(rows)]
    idx = 0
    for c in order:
        for r in range(rows):
            matrix[r][c] = ciphertext[idx] if idx < len(ciphertext) else " "
            idx += 1
    return "".join("".join(row) for row in matrix).rstrip()""")

    add_body_paragraph(doc, "Перевірка на прикладі з методичних вказівок: відкритий текст «Прикладна математика» (20 символів), ключ «Шифр». Алфавітні номери літер ключа «Ш и ф р» відповідають порядку [4, 1, 3, 2]. Зчитування стовпців у порядку 1 («и»): «ра еи», 2 («р»): «кнааа», 3 («ф»): «идммк», 4 («Ш»): «платт» дає криптограму: «ра еикнаааидммкплатт» (рис. 4).")

    add_screenshot_placeholder(
        doc,
        "Перейдіть у підвкладку «1.3. Шифр Стовпцевої Перестановки (Матриця)». Натисніть кнопку «[Вперед]», щоб підсвітився один зі стовпців матриці, та зробіть скріншот таблиці перестановки і покрокового логу зчитування",
        "Рисунок 4 - Покрокова візуалізація матриці шифру перестановки за ключем «Шифр»"
    )

    add_heading_line(doc, "Відповіді на контрольні запитання:")
    add_body_paragraph(doc, "Шифр Цезаря - це симетричний алгоритм моноалфавітної підстановки, в якому кожен символ відкритого тексту замінюється символом, що знаходиться на фіксованій кількості позицій K далі в алфавіті (у класичному варіанті K = 3). Досягнувши кінця алфавіту, відлік циклічно переходить на його початок за правилом модульної арифметики C = (M + K) mod N.", bold_prefix="1. Опишіть алгоритм шифрування Цезаря. ")
    add_body_paragraph(doc, "Метод базується на тому, що в будь-якій природній мові окремі літери зустрічаються з відомою середньостатистичною частотою. Оскільки при моноалфавітній заміні однакова літера завжди переходить в один і той самий символ криптограми, криптоаналітик знаходить найчастіший символ у шифротексті, прирівнює його до найчастішої літери мови (наприклад, «О» для української або «E» для англійської) і знаходить ключ K як різницю їхніх індексів в алфавіті.", bold_prefix="2. У чому суть методу частотного криптоаналізу? ")

    add_heading_line(doc, "Висновок:")
    add_body_paragraph(doc, "Під час виконання лабораторної роботи було детально вивчено принципи роботи класичних шифрів моноалфавітної заміни та стовпцевої перестановки. Практично розроблено програмне забезпечення з покроковою візуалізацією процесу шифрування та дешифрування для українського й англійського алфавітів. За допомогою реалізованого спектрометра частотного аналізу було досліджено вразливість моноалфавітних шифрів до статистичного криптоаналізу та автоматично дешифровано тестову криптограму шляхом визначення ключа K = 10. У третій частині роботи було перевірено матричний алгоритм стовпцевої перестановки, який зберігає природні частоти літер, проте змінює їхній порядок відповідно до рангу символів ключового слова.")

    out_path = os.path.join(OUT_DIR, "Zvit_Lab_1.docx")
    doc.save(out_path)
    print("Saved:", out_path)


def build_lab2():
    doc = prepare_base_doc(
        lab_num=2,
        course_name="Криптографія та захист інформації",
        topic_name="АВТЕНТИФІКАЦІЯ З ВИКОРИСТАННЯМ СТОРОННІХ СЕРВІСІВ. JSON ВЕБ ТОКЕНИ (JWT)"
    )

    add_meta_paragraph(
        doc,
        "ознайомитися з механізмами безсесійної (stateless) автентифікації у вебдодатках за допомогою стандарту JSON Web Token (JWT) та протоколу OAuth 2.0, дослідити внутрішню структуру токена (Header, Payload, Signature), а також практично реалізувати вебсервер із хешуванням паролів за алгоритмом SHA-256 та захищеними маршрутами на основі цифрового підпису HMAC-SHA256 (HS256)."
    )

    add_heading_line(doc, "Завдання до роботи:", centered=False)
    add_body_paragraph(doc, "Індивідуальні вхідні дані для студента №12 у списку групи КБ-305 (Марчук Дмитро Андрійович):")
    add_body_paragraph(doc, "- Користувач 1 (на 1 вище у списку, id = 11): Максим Марія Романівна (логін: «Максим», пароль: «Марія»).")
    add_body_paragraph(doc, "- Користувач 2 (поточний студент, id = 12): Марчук Дмитро Андрійович (логін: «Марчук», пароль: «Дмитро»).")
    add_body_paragraph(doc, "- Користувач 3 (на 1 нижче у списку, id = 13): Новак Давід (логін: «Новак», пароль: «Давід»).")
    add_body_paragraph(doc, "- Алгоритм зберігання та перевірки паролів у базі даних: простий хеш SHA-256 без солі.")
    add_body_paragraph(doc, "- Секретний ключ підпису JWT (Secret): у форматі FirstnameLastname - «DmytroMarchuk» (незалежно від користувача). Час дії токена: 12.")

    add_heading_line(doc, "Хід роботи", centered=True)

    add_heading_line(doc, "1. Налаштування вебсервера та бази даних користувачів із хешуванням SHA-256")
    add_body_paragraph(doc, "Для виконання завдання було розроблено серверний скрипт на мові Python з використанням фреймворку Flask, бібліотек hashlib та PyJWT, а також інтерактивний веб-модуль у складі комплексу AEGIS CRYPTO-OPS (вкладка «ЛАБ 02»). Відповідно до вимог завдання, паролі користувачів (їхні імена) не зберігаються у відкритому вигляді, а хешуються за допомогою криптографічної функції SHA-256 без солі:")
    add_body_paragraph(doc, "- для користувача «Максим» (id = 11, пароль «Марія»): SHA-256 = b3b4d22f55ef859782bc3f175083045b18c1e747eb037f2506d2c36b99c49661;")
    add_body_paragraph(doc, "- для користувача «Марчук» (id = 12, пароль «Дмитро»): SHA-256 = 64c694e98c5f08aa3b805eb9bf4d7be11e409d780388b0ce7020fd09e943a91e;")
    add_body_paragraph(doc, "- для користувача «Новак» (id = 13, пароль «Давід»): SHA-256 = 00f8230510f863b528096443d3efbe208f7ab70733991e7a203e7bc075e6ef95.")
    add_body_paragraph(doc, "Програмний код вебсервера авторизації та перевірки JWT (файл lab2_jwt_server.py):")

    add_code_block(doc, """import hashlib
import datetime
import jwt
from flask import Flask, request, jsonify

app = Flask(__name__)

JWT_SECRET = "DmytroMarchuk"
GROUP_LIST_NUMBER = 12
TOKEN_TTL_SECONDS = GROUP_LIST_NUMBER * 60

def sha256_hash(password: str) -> str:
    return hashlib.sha256(password.encode("utf-8")).hexdigest()

USERS_TRIAD = [
    {"id": 11, "login": "Максим", "password_hash": sha256_hash("Марія")},
    {"id": 12, "login": "Марчук", "password_hash": sha256_hash("Дмитро")},
    {"id": 13, "login": "Новак",  "password_hash": sha256_hash("Давід")},
]

@app.route("/login", methods=["POST"])
def login():
    data = request.get_json(silent=True) or {}
    login_input = data.get("login", "")
    password_input = data.get("password", "")

    input_hash = sha256_hash(password_input)
    user = next((u for u in USERS_TRIAD if u["login"] == login_input and u["password_hash"] == input_hash), None)
    if not user:
        return jsonify({"error": "Невірний логін або пароль", "input_sha256": input_hash}), 401

    payload = {
        "id": user["id"],
        "login": user["login"],
        "exp": datetime.datetime.now(datetime.timezone.utc) + datetime.timedelta(seconds=TOKEN_TTL_SECONDS)
    }
    token = jwt.encode(payload, JWT_SECRET, algorithm="HS256")
    return jsonify({
        "message": f"Користувач {user['login']} (id={user['id']}) був успішно авторизований!",
        "token": token
    }), 200

@app.route("/protected", methods=["GET"])
def protected_route():
    auth_header = request.headers.get("Authorization", "")
    if not auth_header.startswith("Bearer "):
        return jsonify({"error": "Відсутній заголовок Authorization: Bearer <token>"}), 401

    token = auth_header.split(" ", 1)[1]
    try:
        decoded = jwt.decode(token, JWT_SECRET, algorithms=["HS256"])
        return jsonify({"status": "Доступ дозволено!", "user_data": decoded}), 200
    except jwt.ExpiredSignatureError:
        return jsonify({"error": "Час дії токена минув!"}), 401
    except jwt.InvalidTokenError as e:
        return jsonify({"error": f"Невалідний токен: {str(e)}"}), 401""")

    add_heading_line(doc, "2. Перевірка авторизації на маршруті /login та доступу за Bearer-токеном")
    add_body_paragraph(doc, "Під час звернення методом POST на шлях http://localhost:5000/login з обліковими даними студента №12 (login: «Марчук», password: «Дмитро») сервер обчислює SHA-256 хеш введеного пароля, звіряє його із записом у базі даних та повертає повідомлення про успішну авторизацію разом із підписаним JWT-токеном (рис. 1).")

    add_screenshot_placeholder(
        doc,
        "Відкрийте сайт http://localhost:8090 -> вкладка «ЛАБ 02» -> підвкладка «2.1. База Користувачів (SHA-256) та Симулятор /login + JWT». Зробіть скріншот, де видно таблицю трійки користувачів (11. Максим, 12. Марчук, 13. Новак) та вікно Cyber-Postman зі статусом «[HTTP 200 OK]: Користувач Марчук (id=12) був успішно авторизований!»",
        "Рисунок 1 - Таблиця користувачів групи з хешами SHA-256 та успішна авторизація на шляху /login"
    )

    add_body_paragraph(doc, "Також було протестовано відображення повного списку групи з 25 осіб та спробу входу з невірним паролем (коли обчислений хеш SHA-256 не збігається з еталонним хешем у базі даних, що призводить до відповіді HTTP 401 Unauthorized) (рис. 2).")

    add_screenshot_placeholder(
        doc,
        "У тій самій вкладці «2.1» перемкніть кнопку «Увесь список групи (25 осіб)», а в полі пароля введіть будь-яке невірне слово і натисніть «Send Request», щоб показати помилку HTTP 401 Unauthorized",
        "Рисунок 2 - Повний список користувачів групи та блокування входу при невірному хеші пароля"
    )

    add_heading_line(doc, "3. Деконструкція структури JWT та перевірка стійкості підпису до підробки")
    add_body_paragraph(doc, "Згенерований токен складається з трьох Base64Url-кодованих частин (Header.Payload.Signature): Заголовок (Header) містить {\"alg\": \"HS256\", \"typ\": \"JWT\"}; Корисне навантаження (Payload) містить \"id\": 12, \"login\": \"Марчук\", \"iat\" та \"exp\"; Підпис (Signature) обчислюється як HMACSHA256(base64UrlEncode(header) + \".\" + base64UrlEncode(payload), \"DmytroMarchuk\") (рис. 3).")

    add_screenshot_placeholder(
        doc,
        "Перейдіть у підвкладку «2.2. Візуальний Деконструктор JWT (jwt.io) та Лабораторія Підробки». Зробіть скріншот розкладеного на три кольорові частини токена зі статусом «[ПІДПИС ВЕРИФІКОВАНО]»",
        "Рисунок 3 - Покрокова деконструкція та верифікація JWT-токена"
    )

    add_body_paragraph(doc, "Для демонстрації криптографічного захисту цілісності токена було виконано симуляцію атаки підміни корисного навантаження (зміна поля id з 12 на 999 та login на HACKER_ADMIN без знання секретного ключа DmytroMarchuk). Під час перевірки модифікованого токена сервер перераховує HMAC-SHA256 підпис і виявляє невідповідність, відхиляючи запит до захищеного маршруту (рис. 4).")

    add_screenshot_placeholder(
        doc,
        "У підвкладці «2.2» натисніть кнопку «[АТАКА] Симулювати підміну ID у Payload без знання Secret» та зробіть скріншот повідомлення «[ПІДПИС НЕДІЙСНИЙ / ТОКЕН ВІДХИЛЕНО]»",
        "Рисунок 4 - Виявлення спроби несанкціонованої модифікації Payload за рахунок невідповідності підпису HMAC-SHA256"
    )

    add_heading_line(doc, "Відповіді на контрольні запитання:")
    add_body_paragraph(doc, "JWT (JSON Web Token) - це відкритий стандарт (RFC 7519), який визначає компактний та самодостатній спосіб безпечної передачі інформації між сторонами у вигляді JSON-об'єкта, захищеного криптографічним підписом.", bold_prefix="1. Що таке JWT? ")
    add_body_paragraph(doc, "Компактна форма JWT складається з трьох рядків у кодуванні Base64Url, розділених крапками (xxxxx.yyyyy.zzzzz): Заголовок (Header), Корисне навантаження (Payload) та Підпис (Signature).", bold_prefix="2. Структура токена. ")
    add_body_paragraph(doc, "У заголовку зберігаються метадані токена: тип токена (\"typ\": \"JWT\") та назва алгоритму підпису/хешування (\"alg\": \"HS256\"). У корисному навантаженні передаються твердження (claims) про користувача та додаткові параметри (id, login, exp, iat).", bold_prefix="3. Які дані зберігаються в заголовку та корисному навантаженні? ")
    add_body_paragraph(doc, "Підпис формується з трьох компонентів: 1) Base64Url-кодованого заголовка (Header); 2) Base64Url-кодованого навантаження (Payload); 3) секретного ключа сервера (Secret), які передаються у функцію HMACSHA256.", bold_prefix="4. Назвіть частини підпису. ")
    add_body_paragraph(doc, "Користувач передає облікові дані на сервер автентифікації. Сервер перевіряє хеш пароля, формує та підписує JWT своїм секретним ключем і повертає токен клієнту. При кожному зверненні до захищених маршрутів API клієнт додає токен у заголовок Authorization: Bearer <token>, а сервер перевіряє валідність підпису та термін дії exp.", bold_prefix="5. Опишіть схему автентифікації з використанням JWT. ")

    add_heading_line(doc, "Висновок:")
    add_body_paragraph(doc, "Під час виконання лабораторної роботи було вивчено архітектуру безсесійної автентифікації на основі JSON Web Tokens (JWT) та схему делегованої авторизації OAuth 2.0. Практично створено базу даних користувачів групи (варіант 12 - Марчук Дмитро та сусідні за списком студенти), у якій паролі захищено хешуванням SHA-256. Реалізовано механізм генерації та верифікації токенів HS256 із секретним ключем DmytroMarchuk та часом дії 12, а також експериментально підтверджено неможливість підробки даних всередині Payload без знання секретного ключа сервера.")

    out_path = os.path.join(OUT_DIR, "Zvit_Lab_2.docx")
    doc.save(out_path)
    print("Saved:", out_path)


def build_lab3():
    doc = prepare_base_doc(
        lab_num=3,
        course_name="Криптографія та захист інформації",
        topic_name="СИСТЕМА БЛОЧНОГО ШИФРУВАННЯ S-DES"
    )

    add_meta_paragraph(
        doc,
        "вивчити архітектуру та математичний апарат симетричного блочного шифрування на прикладі спрощеного алгоритму S-DES (Simplified DES), дослідити процедуру розгортання раундових підключів (K1, K2), роботу початкової та оберненої перестановок (IP, IP^-1), нелінійних таблиць заміни (S0, S1) і раундової функції Мейстера-Фейстеля, а також виконати покрокове дешифрування індивідуальної криптограми."
    )

    add_heading_line(doc, "Завдання до роботи:", centered=False)
    add_body_paragraph(doc, "Використовуючи 10-бітний ключ (0011011101), розшифрувати криптограму свого варіанту, зашифровану алгоритмом S-DES. Продемонструвати проміжні результати, отримані на виході кожної функції (IP, f_K, SW, f_K, IP^-1), та перетворити отриману двійкову послідовність у десяткове число.")
    add_body_paragraph(doc, "Індивідуальні дані для Варіанта 12 (Марчук Дмитро Андрійович): вхідна 8-бітна криптограма C = 01110011, 10-бітний ключ K = 0011011101.")

    add_heading_line(doc, "Хід роботи", centered=True)

    add_heading_line(doc, "1. Покрокова процедура генерування раундових підключів K1 та K2")
    add_body_paragraph(doc, "Вхідний 10-бітний ключ: K = (0011011101).")
    add_body_paragraph(doc, "- Крок 1. Застосовуємо перестановку P10 = [3, 5, 2, 7, 4, 10, 1, 9, 8, 6]: P10(0011011101) = (1001110011).")
    add_body_paragraph(doc, "- Крок 2. Розділяємо отриману послідовність на дві 5-бітні половини: L = 10011, R = 10011. Виконуємо циклічний зсув ліворуч на 1 позицію (LS-1) для кожної половини окремо: LS-1(10011) = 00111, LS-1(10011) = 00111 -> об'єднано: (00111 00111).")
    add_body_paragraph(doc, "- Крок 3. Застосовуємо перестановку стиснення P8 = [6, 3, 7, 4, 8, 5, 10, 9] для отримання першого підключа K1: K1 = P8(0011100111) = (01011111).")
    add_body_paragraph(doc, "- Крок 4. До половин після першого зсуву (00111 та 00111) застосовуємо циклічний зсув ліворуч на 2 позиції (LS-2): LS-2(00111) = 11100, LS-2(00111) = 11100 -> об'єднано: (11100 11100).")
    add_body_paragraph(doc, "- Крок 5. Застосовуємо перестановку P8 для отримання другого підключа K2: K2 = P8(1110011100) = (11101000).")
    add_body_paragraph(doc, "Оскільки виконується розшифрування, у першому раунді f_K використовується підключ K2 = 11101000, а у другому раунді - підключ K1 = 01011111 (рис. 1).")

    add_screenshot_placeholder(
        doc,
        "Відкрийте сайт http://localhost:8090 -> вкладка «ЛАБ 03» -> підвкладка «3.1. Покроковий Бітовий Візуалізатор S-DES». Переконайтеся, що вибрано Варіант 12 (01110011), і зробіть скріншот верхньої частини сторінки зі стрічкою всіх проміжних результатів та схемою комутації бітів P10/P8",
        "Рисунок 1 - Інтерфейс трасувальника S-DES для Варіанта 12 та схема комутації бітів підключів"
    )

    add_heading_line(doc, "2. Покрокове розшифрування криптограми Варіанта 12 (01110011)")
    add_body_paragraph(doc, "До вхідної криптограми C = (01110011) застосовуємо таблицю IP = [2, 6, 3, 1, 4, 8, 5, 7]: IP(01110011) = 10101101. Розбиваємо блок на ліву та праву 4-бітні половини: L0 = 1010, R0 = 1101.", bold_prefix="Операція 1: Початкова перестановка IP. ")
    add_body_paragraph(doc, "Виконуємо розширення з перестановкою E/P = [4, 1, 2, 3, 2, 3, 4, 1] над правою половиною R0 = 1101: E/P(1101) = 11101011. Додаємо за модулем 2 (XOR) із підключем K2 = 11101000: 11101011 XOR 11101000 = 00000011. Перші 4 біти (0000) подаємо на матрицю S0: рядок = 00 (0), стовпець = 00 (0) -> S0[0][0] = 1 = 01. Другі 4 біти (0011) подаємо на матрицю S1: рядок = 01 (1), стовпець = 01 (1) -> S1[1][1] = 0 = 00. Об'єднаний вихід S-блоків: 0100. Пряма перестановка P4 = [2, 4, 3, 1] дає P4(0100) = 1000. Побітовий XOR із лівою половиною L0 = 1010: 1010 XOR 1000 = 0010. Права половина R0 = 1101 залишається без змін. Отже, вихід першої функції f_K(K2) дорівнює 00101101 (рис. 2).", bold_prefix="Операція 2: Перший раунд циклової функції f_K з підключем K2 = 11101000. ")

    add_screenshot_placeholder(
        doc,
        "У вкладці «3.1» натисніть кнопку «Раунд 1 (f_K)» над матрицями S0 та S1 і зробіть скріншот перехресного підсвічування комірок у матрицях заміни S0 та S1",
        "Рисунок 2 - Робота модулів заміни S0 та S1 у першому раунді дешифрування (підключ K2)"
    )

    add_body_paragraph(doc, "Міняємо місцями ліві та праві 4 біти результату першого раунду (0010 | 1101): SW(00101101) = 11010010. Тепер L1 = 1101, R1 = 0010.", bold_prefix="Операція 3: Перестановка половин блоку SW (Switch). ")
    add_body_paragraph(doc, "Розширення з перестановкою E/P над R1 = 0010 дає E/P(0010) = 00010100. Побітовий XOR із підключем K1 = 01011111: 00010100 XOR 01011111 = 01001011. Перші 4 біти (0100) у S0: рядок = 00 (0), стовпець = 10 (2) -> S0[0][2] = 3 = 11. Другі 4 біти (1011) у S1: рядок = 11 (3), стовпець = 01 (1) -> S1[3][1] = 1 = 01. Об'єднаний вихід S-блоків: 1101. Перестановка P4(1101) = 1101. Побітовий XOR із лівою половиною L1 = 1101: 1101 XOR 1101 = 0000. Права половина R1 = 0010 передається без змін. Вихід другої функції f_K(K1) дорівнює 00000010 (рис. 3).", bold_prefix="Операція 4: Другий раунд циклової функції f_K з підключем K1 = 01011111. ")

    add_screenshot_placeholder(
        doc,
        "Натисніть кнопку «Раунд 2 (f_K)» та зробіть скріншот матриць S0 та S1 для другого раунду дешифрування (підключ K1) і повного покрокового протоколу обчислень унизу",
        "Рисунок 3 - Робота модулів заміни S0 та S1 у другому раунді та повний протокол обчислень"
    )

    add_body_paragraph(doc, "Застосовуємо таблицю IP^-1 = [4, 1, 3, 5, 7, 2, 8, 6] до блоку (00000010): IP^-1(00000010) = 00001000. Перетворимо отриману двійкову послідовність M = 00001000 у десяткове число: 0*2^7 + 0*2^6 + 0*2^5 + 0*2^4 + 1*2^3 + 0*2^2 + 0*2^1 + 0*2^0 = 8 (рис. 4).", bold_prefix="Операція 5: Кінцева обернена перестановка IP^-1 та переведення у десяткове число. ")

    add_screenshot_placeholder(
        doc,
        "Перейдіть у підвкладку «3.2. Зведена Таблиця Розв'язків Усіх 25 Варіантів» і зробіть скріншот таблиці, де підсвічено рядок «Варіант 12 - Марчук Дмитро Андрійович»",
        "Рисунок 4 - Зведена таблиця результатів дешифрування S-DES з виділеним Варіантом 12"
    )

    add_heading_line(doc, "3. Програмна реалізація алгоритму S-DES на мові Python:")
    add_code_block(doc, """P10 = [3, 5, 2, 7, 4, 10, 1, 9, 8, 6]
P8  = [6, 3, 7, 4, 8, 5, 10, 9]
IP  = [2, 6, 3, 1, 4, 8, 5, 7]
IP_INV = [4, 1, 3, 5, 7, 2, 8, 6]
EP  = [4, 1, 2, 3, 2, 3, 4, 1]
P4  = [2, 4, 3, 1]

S0 = [[1, 0, 3, 2], [3, 2, 1, 0], [0, 2, 1, 3], [3, 1, 3, 1]]
S1 = [[1, 1, 2, 3], [2, 0, 1, 3], [3, 0, 1, 0], [2, 1, 0, 3]]

def permute(bits: str, table: list) -> str:
    return "".join(bits[i - 1] for i in table)

def ls(bits: str, n: int) -> str:
    return bits[n:] + bits[:n]

def xor_bits(a: str, b: str) -> str:
    return "".join("0" if x == y else "1" for x, y in zip(a, b))

def sbox_lookup(bits4: str, sbox: list) -> str:
    row = int(bits4[0] + bits4[3], 2)
    col = int(bits4[1] + bits4[2], 2)
    return format(sbox[row][col], "02b")

def generate_subkeys(key10: str):
    p10 = permute(key10, P10)
    l1, r1 = ls(p10[:5], 1), ls(p10[5:], 1)
    k1 = permute(l1 + r1, P8)
    l2, r2 = ls(l1, 2), ls(r1, 2)
    k2 = permute(l2 + r2, P8)
    return k1, k2

def fk(bits8: str, subkey: str) -> str:
    L, R = bits8[:4], bits8[4:]
    xored = xor_bits(permute(R, EP), subkey)
    s_out = sbox_lookup(xored[:4], S0) + sbox_lookup(xored[4:], S1)
    return xor_bits(L, permute(s_out, P4)) + R""")

    add_heading_line(doc, "Висновок:")
    add_body_paragraph(doc, "Під час виконання лабораторної роботи було детально досліджено структуру симетричного блочного шифру S-DES. Виконано покрокове генерування раундових підключів K1 = 01011111 та K2 = 11101000 із заданого 10-бітного ключа (0011011101). Проведено повне трасування процесу розшифрування криптограми Варіанта 12 (01110011) через усі п'ять етапів алгоритму (IP -> f_K(K2) -> SW -> f_K(K1) -> IP^-1), у результаті чого отримано відкритий блок 00001000, що у десятковій системі числення дорівнює 8.")

    out_path = os.path.join(OUT_DIR, "Zvit_Lab_3.docx")
    doc.save(out_path)
    print("Saved:", out_path)


def build_lab4():
    doc = prepare_base_doc(
        lab_num=4,
        course_name="Криптографія та захист інформації",
        topic_name="ПРОТОКОЛ ДІФФІ-ХЕЛЛМАНА ТА АТАКА “ЛЮДИНА ПОСЕРЕДИНІ” (MITM)"
    )

    add_meta_paragraph(
        doc,
        "вивчити математичні основи та схему безпечного обміну криптографічними ключами по відкритому каналу зв'язку за протоколом Діффі-Хеллмана (DH), дослідити вразливість базового протоколу без взаємної автентифікації вузлів до атаки «Людина посередині» (Man-in-the-Middle, MITM), а також практично відтворити роботу протоколу та перехоплення/модифікацію повідомлень у багатовіконному мережевому симуляторі."
    )

    add_heading_line(doc, "Завдання до роботи:", centered=False)
    add_body_paragraph(doc, "1. Локально відтворити протокол Діффі-Хеллмана на 2 учасників (Аліса та Боб).")
    add_body_paragraph(doc, "2. Локально відтворити атаку «Людина посередині» (MITM) на протоколі Діффі-Хеллмана.")
    add_body_paragraph(doc, "3. Провести наочну демонстрацію роботи протоколу та атаки перехоплення і підміни повідомлень між месенджерами учасників.")
    add_body_paragraph(doc, "4. Сформувати звіт (програмний код, математичні розрахунки, скріншоти).")

    add_heading_line(doc, "Хід роботи", centered=True)

    add_heading_line(doc, "1. Відтворення класичного протоколу Діффі-Хеллмана на 2 учасників")
    add_body_paragraph(doc, "Нехай загальнодоступними параметрами, узгодженими довіреною стороною (Trust Party), є просте число p = 331 та генератор мультиплікативної групи g = 204. Покроковий алгоритм формування спільного секретного ключа між Алісою та Бобом:")
    add_body_paragraph(doc, "- Крок 1: Аліса та Боб обирають власні таємні числа в межах [1, p - 1]: Аліса обирає a = 46; Боб обирає b = 56. На цьому етапі Алісі відомо (p=331, g=204, a=46), Бобу відомо (p=331, g=204, b=56).")
    add_body_paragraph(doc, "- Крок 2: Аліса обчислює свій відкритий ключ A та передає його Бобу через відкритий канал: A = g^a mod p = 204^46 mod 331 = 130.")
    add_body_paragraph(doc, "- Крок 3: Боб обчислює свій відкритий ключ B та передає його Алісі: B = g^b mod p = 204^56 mod 331 = 296.")
    add_body_paragraph(doc, "- Крок 4: Обидві сторони незалежно обчислюють спільний секретний ключ K: у Аліси K_A = B^a mod p = 296^46 mod 331 = 58; у Боба K_B = A^b mod p = 130^56 mod 331 = 58. Обидва абоненти отримують ідентичне значення K = 58, тоді як пасивний перехоплювач, знаючи лише (p=331, g=204, A=130, B=296), не може обчислити K без розв'язання задачі дискретного логарифмування (рис. 1).")

    add_screenshot_placeholder(
        doc,
        "Відкрийте сайт http://localhost:8090 -> вкладка «ЛАБ 04» -> підвкладка «4.1. Симулятор Протоколу Діффі-Хеллмана та Атаки MITM». У режимі «Стандартний DH (2 учасники: Аліса <-> Боб)» зробіть скріншот топології мережі, де видно обчислений спільний ключ K = 58 та картки стану вузлів",
        "Рисунок 1 - Візуалізація класичного протоколу обміну ключами Діффі-Хеллмана (K = 58)"
    )

    add_body_paragraph(doc, "Також у підвкладці «4.2» реалізовано покрокове трасування швидкого бінарного піднесення до степеня за модулем (Square-and-Multiply) для перевірки обчислень 204^46 mod 331 = 130 та 204^56 mod 331 = 296 (рис. 2).")

    add_screenshot_placeholder(
        doc,
        "Перейдіть у підвкладку «4.2. Покрокове Піднесення до Степеня за Модулем» і зробіть скріншот таблиць порозрядного обчислення A, B та K",
        "Рисунок 2 - Таблиці бінарного піднесення до степеня за модулем для ключів Аліси та Боба"
    )

    add_heading_line(doc, "2. Відтворення атаки «Людина посередині» (MITM) на протоколі Діффі-Хеллмана")
    add_body_paragraph(doc, "За відсутності автентифікації відкритих ключів активний зловмисник (Attacker) розташовується між Алісою та Бобом і перехоплює весь трафік. Зловмисник генерує власні секретні ключі: c1 = 33 (для зв'язку з Алісою) та c2 = 71 (для зв'язку з Бобом). Коли Аліса відправляє свій відкритий ключ A1 = 130, Зловмисник перехоплює його і замість нього надсилає Бобу свій підставний ключ A2 = g^c2 mod p = 204^71 mod 331 = 325. Коли Боб у відповідь відправляє свій відкритий ключ B2 = 296, Зловмисник перехоплює його і надсилає Алісі свій підставний ключ B1 = g^c1 mod p = 204^33 mod 331 = 267.")
    add_body_paragraph(doc, "У результаті замість одного спільного ключа утворюються два різні сесійні ключі, якими одночасно володіє Зловмисник: 1) Канал «Аліса <-> Зловмисник»: K1 = B1^a mod p = 267^46 mod 331 = 130^33 mod 331 = 64; 2) Канал «Зловмисник <-> Боб»: K2 = A2^b mod p = 325^56 mod 331 = 296^71 mod 331 = 139 (рис. 3).")

    add_screenshot_placeholder(
        doc,
        "У підвкладці «4.1» перемкніть режим на «Атака MITM (Перехоплення: Аліса <-> Зловмисник <-> Боб)» і зробіть скріншот мережевої схеми з трьома вузлами, де видно підмінені ключі A1->A2, B2->B1 та два скомпрометовані ключі K1 = 64 і K2 = 139",
        "Рисунок 3 - Схема атаки «Людина посередині» з утворенням двох сесійних ключів K1 = 64 та K2 = 139"
    )

    add_body_paragraph(doc, "Програмна реалізація протоколу Діффі-Хеллмана та атаки MITM на мові Python:")
    add_code_block(doc, """def diffie_hellman_standard(p: int = 331, g: int = 204, a: int = 46, b: int = 56):
    A = pow(g, a, p)
    B = pow(g, b, p)
    K_alice = pow(B, a, p)
    K_bob = pow(A, b, p)
    print(f"A = {A}, B = {B} -> Спільний секрет K = {K_alice} (у Боба: {K_bob})")
    return K_alice

def diffie_hellman_mitm(p: int = 331, g: int = 204, a: int = 46, b: int = 56, c1: int = 33, c2: int = 71):
    A1 = pow(g, a, p)
    B1 = pow(g, c1, p)
    A2 = pow(g, c2, p)
    B2 = pow(g, b, p)
    K1_alice = pow(B1, a, p)
    K1_attacker = pow(A1, c1, p)
    K2_bob = pow(A2, b, p)
    K2_attacker = pow(B2, c2, p)
    print(f"MITM: K1 (Аліса<->Атакер) = {K1_alice}, K2 (Атакер<->Боб) = {K2_bob}")""")

    add_heading_line(doc, "3. Практична демонстрація перехоплення та модифікації повідомлень у багатовіконному месенджері")
    add_body_paragraph(doc, "Для демонстрації реальної дії атаки MITM було розроблено багатовіконний веб-месенджер (messenger.html), який синхронізує в реальному часі три окремі вкладки браузера через BroadcastChannel: Термінал Аліси (ключ K1 = 64), Консоль Перехоплювача (ключі K1 = 64 та K2 = 139) і Термінал Боба (ключ K2 = 139).")
    add_body_paragraph(doc, "У вкладці Месенджера Аліси було відправлено повідомлення: «Переказ 500 грн на рахунок 1111». Повідомлення було зашифроване ключем Аліси K1 = 64 та затримане вузлом посередника (рис. 4).")

    add_screenshot_placeholder(
        doc,
        "Відкрийте вкладку Месенджера Аліси (http://localhost:8090/messenger.html?role=alice), надішліть повідомлення «Переказ 500 грн на рахунок 1111» і зробіть скріншот вікна месенджера Аліси зі статусом пакета та логом шифрування ключем K = 64 праворуч",
        "Рисунок 4 - Відправлення зашифрованого повідомлення з термінала Аліси (ключ сесії K1 = 64)"
    )

    add_body_paragraph(doc, "У вкладці Консолі Перехоплювача перехоплений пакет було автоматично розшифровано ключем K1 = 64. Далі текст повідомлення було змінено зловмисником на «Переказ 50000 грн на рахунок 9999», перешифровано другим ключем K2 = 139 і відправлено Бобу (рис. 5).")

    add_screenshot_placeholder(
        doc,
        "Перейдіть у вкладку Консолі Перехоплювача (http://localhost:8090/messenger.html?role=attacker), змініть у полі перехопленого пакета суму/рахунок і зробіть скріншот картки перехоплення та журналу MITM-транзакцій",
        "Рисунок 5 - Консоль перехоплювача MITM під час розшифрування (K1 = 64), модифікації тексту та повторного шифрування (K2 = 139)"
    )

    add_body_paragraph(doc, "У вкладці Месенджера Боба отримана криптограма успішно розшифрувалася його ключем K2 = 139. Боб отримав підмінене повідомлення, не підозрюючи про втручання посередника (рис. 6).")

    add_screenshot_placeholder(
        doc,
        "Перейдіть у вкладку Месенджера Боба (http://localhost:8090/messenger.html?role=bob) після натискання кнопки «Підмінити текст, перешифрувати і відправити» у перехоплювачі та зробіть скріншот отриманого Бобом зміненого повідомлення",
        "Рисунок 6 - Термінал Боба з розшифрованим ключем K2 = 139 підміненим повідомленням"
    )

    add_heading_line(doc, "Відповіді на контрольні запитання:")
    add_body_paragraph(doc, "Протокол Діффі-Хеллмана - це криптографічний протокол асиметричного типу, який дозволяє двом або більше сторонам отримати спільний секретний ключ, обмінюючись інформацією через незахищений публічний канал зв'язку.", bold_prefix="1. Що таке алгоритм Діффі-Хеллмана? ")
    add_body_paragraph(doc, "Сторони узгоджують відкриті параметри: велике просте число p і генератор g. Перший учасник генерує приватне число a і надсилає другому A = g^a mod p. Другий учасник генерує приватне число b і надсилає першому B = g^b mod p. Далі перший обчислює K = B^a mod p, а другий обчислює K = A^b mod p. Обидва отримують однакове число K = g^(ab) mod p.", bold_prefix="2. Опишіть схему обміну ключами за протоколом Діффі-Хеллмана. ")
    add_body_paragraph(doc, "Атака «людина посередині» (MITM) - це вид активної кібератаки, коли зловмисник потайки вбудовується в канал зв'язку між двома абонентами, перехоплює, читає і модифікує повідомлення, видаючи себе за кожну зі сторін. При цьому порушуються автентичність, конфіденційність та цілісність інформації.", bold_prefix="3. Що таке атака «людина посередині» та які властивості інформації вона порушує? ")
    add_body_paragraph(doc, "Зловмисник перехоплює відкритий ключ Аліси A1 і надсилає Бобу свій підставний ключ A2 = g^c2 mod p. Одночасно він перехоплює ключ Боба B2 і надсилає Алісі свій ключ B1 = g^c1 mod p. У результаті встановлюються два окремі ключі: K1 = g^(a*c1) mod p між Алісою і зловмисником та K2 = g^(b*c2) mod p між зловмисником і Бобом.", bold_prefix="4. Опишіть схему обміну ключами при атаці «людина посередині». ")

    add_heading_line(doc, "Висновок:")
    add_body_paragraph(doc, "Під час виконання лабораторної роботи було досліджено математичні принципи протоколу обміну ключами Діффі-Хеллмана та його вразливість до атаки «Людина посередині» (MITM). Практично відтворено роботу класичного протоколу для двох учасників (отримано спільний секретний ключ K = 58 при p = 331, g = 204, a = 46, b = 56). За допомогою розробленого багатовіконного мережевого симулятора продемонстровано перехоплення відкритих ключів зловмисником, утворення двох скомпрометованих сесійних ключів (K1 = 64 та K2 = 139), а також перехоплення, розшифрування, модифікацію та повторне шифрування повідомлень у реальному часі.")

    out_path = os.path.join(OUT_DIR, "Zvit_Lab_4.docx")
    doc.save(out_path)
    print("Saved:", out_path)


if __name__ == "__main__":
    build_lab1()
    build_lab2()
    build_lab3()
    build_lab4()
